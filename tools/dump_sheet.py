"""Full dump of a sheet: values + formulas."""
import sys
import openpyxl
from openpyxl.utils import get_column_letter

SRC = r"C:\Users\thomasjulie\OneDrive - Microsoft\04 - Doc\04 - Licensing\Microsoft_profiling_licencesDyn365_Sample.xlsx"


def dump(sheet, r1, r2, c1, c2, out):
    wbf = openpyxl.load_workbook(SRC, data_only=False)
    wbv = openpyxl.load_workbook(SRC, data_only=True)
    wsf = wbf[sheet]
    wsv = wbv[sheet]
    lines = []
    for r in range(r1, min(r2, wsf.max_row) + 1):
        for c in range(c1, min(c2, wsf.max_column) + 1):
            f = wsf.cell(row=r, column=c).value
            v = wsv.cell(row=r, column=c).value
            if f is None and v is None:
                continue
            ref = f"{get_column_letter(c)}{r}"
            if isinstance(f, str) and f.startswith("="):
                lines.append(f"{ref}\tFORMULA\t{f}\tVAL={v!r}")
            else:
                lines.append(f"{ref}\tVALUE\t{f!r}")
    with open(out, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))
    print(f"{len(lines)} cells -> {out}")


if __name__ == "__main__":
    sheet = sys.argv[1]
    r1, r2, c1, c2 = (int(x) for x in sys.argv[2:6])
    dump(sheet, r1, r2, c1, c2, sys.argv[6])
