"""Extraction brute du classeur Excel source pour retro-engineering."""
import sys
import json
import openpyxl

SRC = r"C:\Users\thomasjulie\OneDrive - Microsoft\04 - Doc\04 - Licensing\Microsoft_profiling_licencesDyn365_Sample.xlsx"


def main():
    wb = openpyxl.load_workbook(SRC, data_only=False)
    print("SHEETS:", wb.sheetnames)
    for ws in wb.worksheets:
        print(f"  - {ws.title}: dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column} state={ws.sheet_state}")
    print("DEFINED NAMES:", list(wb.defined_names.keys()))


if __name__ == "__main__":
    main()
