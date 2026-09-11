"""Dump the structure of the exported workbook to check readability."""
import sys
from openpyxl import load_workbook

path = sys.argv[1] if len(sys.argv) > 1 else "tools/sample-export.xlsx"
wb = load_workbook(path)

print("TABS:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print("\n" + "=" * 78)
    print(f"SHEET: {name}   ({ws.max_row} rows x {ws.max_column} columns)")
    print("=" * 78)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)):
        cells = []
        for c in row:
            if c.value is None:
                continue
            v = str(c.value).replace("\n", " | ")
            if len(v) > 58:
                v = v[:55] + "..."
            cells.append(f"{c.coordinate}={v}")
        if cells:
            print("  " + "   ".join(cells))
