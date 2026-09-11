"""Structured extraction: process rows (B, C, E) + styling (fill colour) to tell domains from processes apart."""
import openpyxl
from openpyxl.utils import get_column_letter

SRC = r"C:\Users\thomasjulie\OneDrive - Microsoft\04 - Doc\04 - Licensing\Microsoft_profiling_licencesDyn365_Sample.xlsx"

wb = openpyxl.load_workbook(SRC, data_only=False)
ws = wb["Profiling_By_Business Process"]

out = []
for r in range(17, 158):
    b = ws.cell(row=r, column=2).value
    c = ws.cell(row=r, column=3).value
    e = ws.cell(row=r, column=5).value
    cell = ws.cell(row=r, column=2)
    fill = cell.fill
    rgb = None
    if fill and fill.fgColor is not None:
        rgb = fill.fgColor.rgb if fill.fgColor.type == "rgb" else f"theme{fill.fgColor.theme}/tint{round(fill.fgColor.tint,3)}"
    indent = cell.alignment.indent
    bold = cell.font.bold
    marks = []
    for col in range(8, 61):
        v = ws.cell(row=r, column=col).value
        if v is not None and not (isinstance(v, str) and v.startswith("=")):
            marks.append(f"{get_column_letter(col)}={v}")
    out.append(f"R{r}\tfill={rgb}\tbold={bold}\tind={indent}\tB={b!r}\tE={e!r}\tC={c!r}\tMARKS={' '.join(marks)}")

with open("tools/out_rows.txt", "w", encoding="utf-8") as fh:
    fh.write("\n".join(out))
print(len(out))
